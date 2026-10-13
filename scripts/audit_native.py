"""Inspect original BI containers and compare embedded schemas to project sources.

This executes the Tableau Hyper engine for extract reads, decodes PBIX tables,
and reads XLSX values/formulas. It does NOT execute DAX or render Desktop views.
Run after installing verification/native/requirements.txt.
"""
from pathlib import Path
import hashlib,io,json,sys,tempfile,zipfile,xml.etree.ElementTree as ET
import pandas as pd
from pbixray import PBIXRay
from tableauhyperapi import Connection,HyperProcess,Telemetry
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'verification/native'

def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def records(df):return json.loads(df.to_json(orient='records',date_format='iso'))

def audit():
 report=[]
 for project in sorted(ROOT.glob('0[1-5]_*')):
  sources={p.stem:list(pd.read_csv(p,nrows=0).columns) for p in (project/'Data/clean').glob('*.csv')}
  dest=OUT/project.name;dest.mkdir(parents=True,exist_ok=True)
  for path in sorted((project/'Dashboard').glob('*')):
   if path.suffix not in ('.pbix','.twbx'):continue
   item={'file':str(path.relative_to(ROOT)),'sha256':sha(path),'native_desktop_refresh':'NOT_RUN','native_visual_render':'NOT_RUN','embedded_tables':[],'canonical_source_schemas':sources}
   with zipfile.ZipFile(path) as z:
    item['zip_crc_pass']=z.testzip() is None
    if path.suffix=='.pbix':
     m=PBIXRay(str(path))
     for attribute in ['schema','dax_measures','power_query','relationships']:
      try:
       frame=getattr(m,attribute);frame.to_csv(dest/(attribute+'.csv'),index=False)
      except Exception as error:item[attribute+'_error']=str(error)
     for table in m.tables:
      if str(table).startswith(('LocalDateTable','DateTableTemplate')):continue
      try:
       frame=m.get_table(str(table));item['embedded_tables'].append({'table':str(table),'rows':len(frame),'columns':list(frame.columns)})
       # Small totals are descriptive of embedded data, not DAX measure execution.
       totals={c:float(pd.to_numeric(frame[c],errors='coerce').sum()) for c in frame.columns if c.lower() in ['revenue_usd','revenue','gmv','order_value','quantity','units_sold']}
       if totals:item['embedded_tables'][-1]['direct_column_sums']=totals
      except Exception as error:item['embedded_tables'].append({'table':str(table),'decode_error':str(error)})
     layout=json.loads(z.read('Report/Layout').decode('utf-16-le'))
     item['report_pages']=[{'name':p.get('displayName'),'visual_count':len(p.get('visualContainers',[]))} for p in layout.get('sections',[])]
     item['dax_execution']='NOT_RUN: PBIXRay reads expressions; it is not the Analysis Services DAX engine.'
    else:
     twb=ET.fromstring(z.read(next(n for n in z.namelist() if n.endswith('.twb'))))
     item['tableau_version']=twb.get('version')
     item['declared_fields']=sorted(set(c.get('name','').strip('[]') for c in twb.findall('.//column') if c.get('datatype') and c.get('name')))
     item['worksheets']=[s.get('name') for s in twb.findall('./worksheets/worksheet')]
     item['connection_types']=sorted(set(c.get('class','') for c in twb.findall('.//connection')))
     calculations={c.get('caption') or c.get('name'):c.find('calculation').get('formula') for c in twb.findall('.//column') if c.find('calculation') is not None}
     (dest/'tableau_calculations.json').write_text(json.dumps(calculations,indent=2)+'\n')
     item['embedded_csvs']=[]
     for name in z.namelist():
      if name.endswith('.csv'):
       frame=pd.read_csv(io.BytesIO(z.read(name)));item['embedded_csvs'].append({'file':name,'rows':len(frame),'columns':list(frame.columns)})
     with tempfile.TemporaryDirectory(prefix='tableau-audit-') as temp:
      for name in z.namelist():
       if not name.endswith('.hyper'):continue
       extracted=Path(temp)/'extract.hyper';extracted.write_bytes(z.read(name))
       try:
        with HyperProcess(Telemetry.DO_NOT_SEND_USAGE_DATA_TO_TABLEAU,parameters={'log_dir':temp,'use_tcp_port':'auto'}) as h:
         with Connection(h.endpoint,str(extracted)) as c:
          for schema in c.catalog.get_schema_names():
           for table in c.catalog.get_table_names(schema):
            definition=c.catalog.get_table_definition(table)
            item['embedded_tables'].append({'table':str(table),'rows':c.execute_scalar_query(f'SELECT COUNT(*) FROM {table}'),'columns':[str(col.name).strip('"') for col in definition.columns]})
        item['hyper_engine_read']='PASS'
       except Exception as error:item['hyper_engine_read']='FAILED';item['hyper_error']=str(error)
   item['canonical_schema_matches']={name:any(set(cols).issubset(set(t.get('columns',[]))) for t in item['embedded_tables']+item.get('embedded_csvs',[])) for name,cols in sources.items()}
   item['source_reconciliation']='MISMATCH' if not all(item['canonical_schema_matches'].values()) else 'SCHEMA_ONLY_MATCH_REQUIRES_VALUE_CHECK'
   report.append(item)
  path=next((project/'Excel').glob('*.xlsx')); wb=load_workbook(path,data_only=False,read_only=True);cached=load_workbook(path,data_only=True,read_only=True)
  cells=[]
  for row in wb['Summary']:
   for cell in row:
    if cell.data_type=='f':cells.append({'cell':cell.coordinate,'formula':cell.value,'cached_value':cached['Summary'][cell.coordinate].value})
  excel={'file':str(path.relative_to(ROOT)),'sha256':sha(path),'formula_cells':cells,'source_checks':[],'engine':'OOXML inspection; application recalculation recorded separately'}
  for name,cols in sources.items():
   if name not in wb.sheetnames:excel['source_checks'].append({'table':name,'pass':False,'reason':'missing sheet'});continue
   frame=pd.DataFrame(wb[name].values);frame.columns=frame.iloc[0];frame=frame.iloc[1:].reset_index(drop=True)
   ref=pd.read_csv(project/'Data/clean'/f'{name}.csv')
   equivalent=len(frame)==len(ref) and list(frame.columns)==list(ref.columns)
   if equivalent:
    for col in ref:
     a=frame[col].reset_index(drop=True);b=ref[col]
     if pd.api.types.is_numeric_dtype(b) or pd.api.types.is_bool_dtype(b):
      import numpy as np
      try:equivalent=equivalent and bool(np.allclose(pd.to_numeric(a).astype(float),b.astype(float),rtol=1e-10,atol=1e-10,equal_nan=True))
      except (ValueError,TypeError):equivalent=False
     else:equivalent=equivalent and a.astype(str).eq(b.astype(str)).all()
   excel['source_checks'].append({'table':name,'rows':len(frame),'pass':bool(equivalent)})
  wb.close();cached.close();report.append(excel)
 (OUT/'workbook_audit.json').write_text(json.dumps(report,indent=2,default=str)+'\n')
 for item in report:print(item['file'],item.get('source_reconciliation',all(x['pass'] for x in item.get('source_checks',[]))))
 return report
if __name__=='__main__':audit()
